from openpyxl import workbook
from _peewee_orm import *
from _imports import *
from _openpyxl import *
from datetime import datetime,date, timedelta
from calendar import monthrange
from rich.progress import track
import os

class Model(BaseModel):
    class Meta:
        database = Connection('db.sqlite3', pragmas={'foreign_keys': 1})

class Employee(Model):

    id = PrimaryKeyColumn()
    name = StringColumn()
    pis = StringColumn(max_length=11, unique=True)
    created_at = TimestampColumn()
    updated_at = TimestampColumn()

    class Meta:
        table_name = "employees"

    def time_clock_marking_by_month(self, year, month, day = 1):
        return TimeClockMarking.select().where(TimeClockMarking.employee == self,
                                               TimeClockMarking.date >= datetime(year, month, day).__str__(),
                                               TimeClockMarking.date <= datetime(year, month, monthrange(year, month)[1]).__str__())

    def time_clock_marking_by_period(self, start_date, final_date):
        return TimeClockMarking.select().where(TimeClockMarking.employee == self,
                                               TimeClockMarking.date >= datetime(start_date['year'], start_date['month'], start_date['day']).__str__(),
                                               TimeClockMarking.date <= datetime(final_date['year'], final_date['month'], final_date['day']).__str__())

    def __str__(self) -> str:
        return "ID: "+str(self.id)+" NAME: "+self.name+" PIS: "+self.pis+" CREATED_AT: "+self.created_at.__str__()+" UPDATED_AT: "+self.updated_at.__str__()
    
    def to_list(self):
        return [self.id, self.name, self.pis, f"{self.created_at}", f"{self.updated_at}"]
    
    @staticmethod
    def all_to_list():
        employees = []

        for e in Employee.select():
            employees.append(e.to_list())
        
        return employees

class TimeClockMarking(Model):
    
    date = DateColumn()
    first_entry = TimeColumn(null=True)
    first_exit = TimeColumn(null=True)
    second_entry = TimeColumn(null=True)
    second_exit = TimeColumn(null=True)
    pis = StringColumn(11)
    employee = ForeignKeyColumn(Employee, on_delete='CASCADE', backref='time_clock_markings')
    
    class Meta:
        table_name = "time_clock_markings"
    
    def __str__(self):
        return "DATE: {}  MARKINGS: {} {} {} {} PIS: {}".format(
            self.date,
            '        ' if self.first_entry is None else self.first_entry,
            '        ' if self.first_exit is None else self.first_exit,
            '        ' if self.second_entry is None else self.second_entry,
            '        ' if self.second_exit is None else self.second_exit,
            self.pis)

class KeyValue(Model):
    
    id = PrimaryKeyColumn()
    key = StringColumn(unique=True)
    value = StringColumn()

    AFD_LINES_READED_COUNT_KEY = 'AFD_LINES_READED_COUNT'
    AFD_FILE_PATH = 'DEFAULT_AFD_FILE_PATH'
    DEFAULT_SPREADSHEET_FOLDER = 'DEFAULT_SPREADSHEET_FOLDER'

    class Meta:
        table_name = "key_values"
    
    def __str__(self):
        return "KEY: "+self.key+"\tVALUE: "+self.value

class Spreadsheet:
    
    #months = {1:'JAN',2:'FEV',3:'MAR',4:'ABR',5:'MAI',6:'JUN',7:'JUL',8:'AGO',9:'SET',10:'OUT',11:'NOV',12:'DEZ',}
    weekdays = {0:'SEG',1:'TER',2:'QUA',3:'QUI',4:'SEX',5:'SAB',6:'DOM'}

    def __init__(self) -> None:
        pass

    def create_workbook(self):
        workbook = Workbook()
        workbook.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
        #workbook.iso_dates = True

        return workbook        

    def get_last_day_of_month(year, month):
        return monthrange(year, month)[1]

    def save_period_db_spreadsheet(self, start_date, final_date, destiny_folder):
        workbook = self.create_workbook()
        markings = workbook.active
        markings.title = 'Marcações'
        markings.freeze_panes = 'M2'
        errors = workbook.create_sheet('Erros')
        errors.freeze_panes = 'M2'
        header = ['NOME', 'DIA','DATA', 'E1', 'S1', 'E2', 'S2', '1ª TURNO', 'ALMOÇO', 'INT.ENTRE.JORNADAS', 'HORA. EXTRA', 'OBS']
        #           A       B     C      D     E     F     G          H        I               J                   K          L

        markings.append(header)
        errors.append(header)

        employees = Employee.select()
        total = employees.__len__()
        line = 2
        line_error = 2

        for index in track(range(0, total), 'Processando...'):
            e = employees[index]
            previous = None

            for t in e.time_clock_marking_by_period(start_date, final_date):
                first_journey = self.calc_time_diff(t.first_entry, t.first_exit)
                lunch = self.calc_time_diff(t.first_exit, t.second_entry)
                break_working = self.calc_break_working_hours(previous, t)
                extra = self.calc_extra_hour(t)
                obs = ''

                row = [e.name, Spreadsheet.weekdays[t.date.weekday()], t.date, t.first_entry, t.first_exit, t.second_entry, t.second_exit, first_journey, lunch, break_working, extra ]
                
                error_journey, obs_journey = self.check_first_journey(first_journey, markings, errors, line, line_error)
                error_lunch, obs_lunch = self.check_lunch(lunch, markings, errors, line, line_error)
                error_break_working, obs_break_working = self.check_break_working(break_working, markings, errors, line, line_error)
                error_extra, obs_extra = self.check_extra(extra, markings, errors, line, line_error)
                
                obs += obs_journey
                obs += obs_lunch
                obs += obs_break_working
                obs += obs_extra

                row.append(obs)
                markings.append(row)
                
                if error_journey or error_lunch or error_break_working or error_extra:
                    errors.append(row)
                    line_error += 1

                previous = t
                line += 1
                    
        markings.add_table(self.create_table('TableStyleMedium9', 'Marcações', line))
        errors.add_table(self.create_table('TableStyleMedium9', 'Erros', line_error))

        if os.name == 'posix': destiny_folder += '/'
        elif os.name == 'nt': destiny_folder += '\\'
        
        file_name = "{}MARCAÇÕES_DE_{}-{}-{}_A_{}-{}-{}.xlsx".format(destiny_folder, start_date['day'], start_date['month'], start_date['year'], final_date['day'], final_date['month'], final_date['year'])
        workbook.save(file_name)

        if os.name == 'nt':
            os.startfile(file_name)

    def check_first_journey(self, first_journey, markings, errors, line, line_error):
        if first_journey != '' and first_journey > timedelta(hours=4, minutes=30):
            #markings[f'H{line}'].font = Font(color="FF0000", italic=True)
            #errors[f'H{line_error}'].font = Font(color="FF0000", italic=True)

            return (True, 'Mais que 4h 30m no primeiro turno, ')
        else:
            return (False, '')

    def check_lunch(self, lunch, markings, errors, line, line_error):
        if lunch != '' and lunch > timedelta(hours=1, minutes=50):
            #markings[f'I{line}'].font = Font(color="FF0000", italic=True)
            #errors[f'I{line_error}'].font = Font(color="FF0000", italic=True)
            
            return (True, 'Mais que 1h 50m de almoço, ')
        else:
            return (False, '')

    def check_break_working(self, break_working, markings, errors, line, line_error):
        if break_working != '' and break_working < timedelta(hours=12):
            #markings[f'J{line}'].font = Font(color="FF0000", italic=True)
            #errors[f'J{line_error}'].font = Font(color="FF0000", italic=True)

            return (True, 'Menos que 12h entre Jornadas, ')
        else:
            return (False, '')

    def check_extra(self, extra, markings, errors, line, line_error):
        if extra != '' and extra > timedelta(minutes=30, hours=1):
            #markings[f'K{line}'].font = Font(color="FF0000", italic=True)
            #errors[f'K{line_error}'].font = Font(color="FF0000", italic=True)

            return (True, 'Mais que 01h 30m extra, ')
        else:
            return (False, '')

    def create_table(self, styleName, displayName, lines):
        table = Table(displayName=displayName, ref=f"A1:L{lines}")
        table.tableStyleInfo = TableStyleInfo(name=styleName, showFirstColumn=True, showLastColumn=True, showRowStripes=True, showColumnStripes=False)

        return table

    def calc_time_diff(self, start_time, end_time):
        if start_time is None or end_time is None: return ''

        d = date(1,1,1)
        start_date = datetime.combine(d, start_time)
        end_date = datetime.combine(d, end_time)

        return end_date - start_date

    def calc_break_working_hours(self, previous, today):
        if previous is None or previous.second_exit is None or today is None or today.first_entry is None:
            return ''
        else:
            start = datetime.combine(previous.date, previous.second_exit)
            end = datetime.combine(today.date, today.first_entry)

            return end - start

    def calc_extra_hour(self, time_clock_marking):
        if time_clock_marking.first_entry is None or time_clock_marking.second_exit is None:
            return ''
        seven_hours_worked = [0,1,2,3]# SEG a QUI
        day_of_week = time_clock_marking.date.weekday()

        d = date(1,1,1)
        start = datetime.combine(d, time_clock_marking.first_entry)
        end = datetime.combine(d, time_clock_marking.second_exit)
        hours_of_lunch = datetime.combine(d, time_clock_marking.second_entry) - datetime.combine(d, time_clock_marking.first_exit)

        worked_hours = (end - start)-hours_of_lunch

        if day_of_week in seven_hours_worked:
            return worked_hours - timedelta(hours=7)
        else:
            return worked_hours - timedelta(hours=8)